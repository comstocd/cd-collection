"""Build collection.js for the CD collection site from CD_Catalogue.xlsx.

Usage: python3 build.py CD_Catalogue.xlsx ../collection.js BUILD

Source sheets:
  Shelf plan       one row per disc case, with section (the master list)
  Picard Jan 2023  MusicBrainz lookups from ripped discs: date, label, catno, barcode, MBID
  Photo batch N    spine transcriptions: where the case was found, label, catno, status notes
Cubby placement reproduces the 4 x 6 shelf plan (rows 1-6 top to bottom, columns 1-4 left to right).
"""
import openpyxl, re, json, sys, collections, datetime

SRC, OUT, BUILD = sys.argv[1], sys.argv[2], int(sys.argv[3])
wb = openpyxl.load_workbook(SRC, read_only=True)

def rows(name):
    r = list(wb[name].iter_rows(values_only=True))
    return [dict(zip(r[0], x)) for x in r[1:]]

def norm(s):
    s = (s or '').lower().replace('’', "'").replace('‐', '-').replace('&', 'and')
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()

# ---------- enrichment pools (consumed in order so duplicate copies get distinct sources) ----------
picard = collections.defaultdict(list)
for r in rows('Picard Jan 2023'):
    picard[(norm(r['artist']), norm(r['title']))].append(r)
photo = collections.defaultdict(list)
for n in wb.sheetnames:
    if not n.startswith('Photo'):
        continue
    for r in rows(n):
        if not r.get('title'):
            continue
        r['_batch'] = n
        photo[(norm(r['artist']), norm(r['title']))].append(r)

# ---------- display cleanup (same rules as the cubby map) ----------
ARTIST_FIX = {
    'Various Artists': 'Various', 'Various (GRP Artists)': 'Various',
    'Art Blakey and The Jazz Messengers': 'Art Blakey & The Jazz Messengers',
    'The Mahavishnu Orchestra with John McLaughlin': 'Mahavishnu Orchestra',
    'Herb Alpert & The Tijuana Brass': "Herb Alpert's Tijuana Brass",
    'Herb Alpert’s Tijuana Brass': "Herb Alpert's Tijuana Brass",
    'Creedence Clearwater Revival featuring John Fogerty': 'Creedence Clearwater Revival',
    'Chicago Transit Authority': 'Chicago',
    'The Jimi Hendrix Experience': 'Jimi Hendrix',
    'Gloria Estefan and Miami Sound Machine': 'Gloria Estefan',
    'Stevie Ray Vaughan and Double Trouble': 'Stevie Ray Vaughan',
    'Al Kooper & Mike Bloomfield': 'Bloomfield & Kooper',
    'Mike Bloomfield, Al Kooper & Stephen Stills': 'Bloomfield & Kooper',
    'Vince Guaraldi Trio': 'Vince Guaraldi',
}
TITLE_FIX = {
    'Music to Your Ears: A Collection of Holiday Music': 'Music to Your Ears',
    'A Charlie Brown Christmas: The Original Sound Track Recording of the CBS Television Special': 'A Charlie Brown Christmas',
    'Glee: The Music, The Christmas Album': 'Glee: The Christmas Album',
    "(I Got No Kick Against) Modern Jazz. A GRP Artists' Celebration Of The Songs Of The Beatles": 'GRP Artists: Songs of the Beatles',
    "A GRP Artists' Celebration of the Songs of the Beatles": 'GRP Artists: Songs of the Beatles',
    'Skeletons From the Closet: The Best of the Grateful Dead': 'Skeletons From the Closet',
    "What a Long Strange Trip It's Been: The Best of the Grateful Dead": "What a Long Strange Trip It's Been",
    'History of the Grateful Dead, Volume 1: Bear’s Choice': "Bear's Choice",
    'Dick’s Picks, Volume 2: Columbus, Ohio 10/31/71': "Dick's Picks Vol. 2",
    'Dick’s Picks, Volume 3: Pembroke Pines, FL 5/22/77': "Dick's Picks Vol. 3",
    'Dick’s Picks, Volume 4: Fillmore East 2/13–14/70': "Dick's Picks Vol. 4",
    'Red Rocks 7/8/78 (3-CD set, handwritten)': 'Red Rocks 7/8/78 (CD-R)',
    'Buffalo, NY, May 9, 1977 (handwritten)': 'Buffalo 5/9/77 (CD-R)',
    'An Evening With the Allman Brothers Band: 2nd Set': '2nd Set',
    '20th Century Masters: The Millennium Collection: The Best of Joe Walsh': 'The Best of Joe Walsh',
    'Studio: Selected Studio Recordings 1986–1995': 'Studio',
    'Fillmore East: The Lost Concert Tapes 12/13/68': 'Fillmore East: Lost Concert Tapes',
    'Viva la Vida or Death and All His Friends': 'Viva la Vida',
    'Paint the Sky with Stars: The Best of Enya': 'Paint the Sky with Stars',
    'The Monty Python Instant Record Collection': 'Instant Record Collection',
    'Master Serie : Édith Piaf': 'Master Serie',
    'Lifescapes: Relaxation for Women': 'Lifescapes',
    'Tapestry Revisited: A Tribute to Carole King': 'Tapestry Revisited (tribute)',
}
CLASSICAL = {  # (filing key, display artist, display title)
    'BBC Music, Volume 3, Number 1: BBC Singers - 70 Years': ('BBC Singers', 'BBC Singers', '70 Years (BBC Music)'),
    'Symphonien nos. 5 & 7': ('Beethoven 1', 'Beethoven', 'Symphonies 5 & 7 (Kleiber)'),
    'BBC Music, Volume 2, Number 7: Piano Sonatas nos. 29, 32': ('Beethoven 2', 'Beethoven', 'Piano Sonatas 29 & 32 (BBC Music)'),
    'BBC Music, Volume 2, Number 8: Brahms: Liebeslieder Waltzes and Neue Liebeslieder / Spohr: Three Psalms, op. 85 / Meyerbeer: Psalm 91': ('Brahms', 'Brahms', 'Liebeslieder Waltzes (BBC Music)'),
    'Chant d’amour : mélodies françaises': ('Bartoli', 'Cecilia Bartoli', "Chant d'amour"),
    'Sentimento': ('Bocelli', 'Andrea Bocelli', 'Sentimento'),
    'Live Recordings, 1957–1963': ('Callas', 'Maria Callas', 'Live Recordings 1957-63'),
    'The Piano Rolls, Volume Two': ('Gershwin', 'Gershwin', 'The Piano Rolls, Vol. 2'),
    'Messiah': ('Handel', 'Handel', 'Messiah'),
    'The Hugo Masters: An Anthology of Chinese Classical Music, Volume 2: Plucked Strings': ('Hugo', 'Various', 'Hugo Masters: Chinese Classical, Vol. 2'),
    'Don Giovanni: Highlights': ('Mozart 1', 'Mozart', 'Don Giovanni highlights'),
    'BBC Music, Volume 2, Number 6: Violin Sonatas in B flat K454, in G K379 & in A K526': ('Mozart 2', 'Mozart', 'Violin Sonatas (BBC Music)'),
    'Vol. 2': ('Pavarotti', 'Luciano Pavarotti', 'Vol. 2'),
    'La bohème': ('Puccini 1', 'Puccini', 'La bohème (Callas)'),
    'Madama Butterfly': ('Puccini 2', 'Puccini', 'Madama Butterfly (Karajan)'),
    'BBC Music, Volume 2, Number 9: Rimsky-Korsakov: The Golden Cockerel, Capriccio Espagnol / Mussorgsky: Pictures at an Exhibition': ('Rimsky', 'Rimsky-Korsakov', 'Golden Cockerel / Pictures (BBC Music)'),
    'BBC Music, Volume 2, Number 10: Liederkreis, op. 39 / Twelve Poems, op. 35': ('Schumann', 'Schumann', 'Liederkreis (BBC Music)'),
    'BBC Music, Volume 2, Number 11: Live From The Proms: Walton: Symphony no. 1 / Takemitsu': ('Walton', 'Walton', 'Symphony No. 1 (BBC Music)'),
}
SOLO = ['Alanis Morissette','Alicia Keys','Alison Moyet','Anita Baker','Aretha Franklin','Bette Midler',
    'Bill Withers','Bob Dylan','Bonnie Raitt','Carole King','Cat Stevens','Celine Dion','Chuck Berry',
    'Corinne Bailey Rae','Diana Ross','Dionne Warwick','Elvis Costello with Burt Bacharach','Eric Clapton',
    'Faith Hill','Frank Zappa','George Benson','Gloria Estefan','Harry Nilsson','Jaci Velasquez',
    'Jackson Browne','James Blunt','Janis Joplin','Jeff Beck','Jennifer Lopez','Jimi Hendrix',
    'Joan Armatrading','Joe Cocker','Joe Walsh','John Mayer','Joni Mitchell','Lyle Lovett','Mariah Carey',
    'Marta Sanchez','Marvin Gaye','Michael Jackson','Miki Howard','Natalie Merchant','Neil Diamond',
    'Neil Young','Ray Charles','Sarah McLachlan','Shawn Colvin','Sheryl Crow','Sinéad O’Connor',
    'Steve Winwood','Stevie Ray Vaughan','Stevie Wonder','Toni Braxton','Tori Amos','Tracy Chapman',
    'Van Morrison','Art Blakey & The Jazz Messengers','Benny Goodman Orchestra','Carlos Munhoz',
    'Ella Fitzgerald','Emanuele Cisi 4et','Harry Connick, Jr.',"Herb Alpert's Tijuana Brass",
    'Herbie Hancock','Illinois Jacquet','Joe Henderson','Keith Jarrett','Louis Jordan','Miles Davis',
    'Norah Jones','Stan Getz','Tony Bennett','Wayne Shorter','Andrea Bocelli','Angelique Kidjo',
    'Amy Hayashi-Jones','Michael Whalen','Luba Tcheresky','Patricia Kaas','Édith Piaf','Jerry Garcia','Mickey Hart']
SURNAME = {}
for a in SOLO:
    w = a.replace('Elvis Costello with Burt Bacharach', 'Elvis Costello').replace(' & The Jazz Messengers', '')
    w = w.replace(' Orchestra', '').replace(' 4et', '').replace(', Jr.', '').replace("'s Tijuana Brass", '')
    last = w.split()[-1]
    if w.startswith('Stevie Ray'): last = 'Vaughan'
    if w == 'Corinne Bailey Rae': last = 'Bailey Rae'
    if w.startswith('Sinéad'): last = 'OConnor'
    if w.startswith('Édith'): last = 'Piaf'
    SURNAME[a] = last + ' ' + w
SURNAME.update({'Bloomfield & Kooper': 'Bloomfield', 'Ozzy Osbourne / Randy Rhoads': 'Osbourne', 'k.d. lang': 'lang', 'Dr. John': 'Dr John', 'Carole King': 'king carole'})

def fkey(a):
    if a in SURNAME: return SURNAME[a].lower()
    k = re.sub(r'^the ', '', a.lower())
    if k.startswith('various') or k.startswith('('): k = 'zzz' + k
    return k

SECS = {'1': ('xmas', 'Christmas'), '2': ('classical', 'Classical'), '3': ('jazz', 'Jazz'), '4': ('beatles', 'Beatles'),
        '5': ('dead', 'Grateful Dead'), '6': ('allmans', 'Allman Brothers'), '7': ('rock', 'Rock & Pop'), '8': ('other', 'Other')}

def copy_info(a0, t0, fmt):
    """Pull pressing details for one physical copy from Picard first, then the photo batches."""
    k = (norm(a0), norm(t0))
    c = {}
    p = picard[k].pop(0) if picard[k] else None
    # a CD and its SACD reissue share artist and title, so match the spine row on format too
    pool = photo[k]
    want = [x for x in pool if ('SACD' in (x['label'] or '') + (x['status'] or '')) == (fmt == 'SACD')]
    ph = want[0] if want else None
    if ph is not None: pool.remove(ph)
    if p:
        c.update(date=str(p['date'] or ''), country=p['country'], label=p['label'], catno=p['catno'],
                 barcode=p['barcode'], mbid=p['mbid'], rtype=p['release_group_type'])
        if 'SACD' in (p['format'] or '') and fmt != 'SACD':
            c['note'] = 'MusicBrainz matched the hybrid SACD edition; check the case for an SACD logo'
            c['sacd_match'] = True
    if ph:
        c.setdefault('label', ph['label'])
        if not c.get('catno'): c['catno'] = ph['catno (from spine)']
        if not c.get('label'): c['label'] = ph['label']
        c['found'] = ph['_batch'].replace('Photo batch', 'Photo') + ', ' + ph['position']
        st = ph['status'] or ''
        if st and not c.get('note') and not st.startswith('In Jan') and st not in ('NEW', 'NEW - SACD', 'Second copy'):
            st = re.sub(r'^NEW( - SACD)?( - |\s)?', '', st).strip()
            c['note'] = st[1:-1] if st.startswith('(') and st.endswith(')') else (st or None)
    c['source'] = 'MusicBrainz (Picard lookup, Jan 2023)' if p else ('Spine photo' if ph else 'Shelf plan')
    return {k2: v for k2, v in c.items() if v not in (None, '')}

# Corrections confirmed against the physical case. Picard's disc-ID lookup matched the wrong
# release for these (a hybrid SACD edition); the barcode on the case is a standard CD.
OVERRIDES = {
    ('Norah Jones', 'Come Away With Me'): dict(discs=1, copy=dict(
        label='Blue Note', barcode='724353208820',
        source='Barcode on the case (checked against library catalogue records for this UPC)')),
    ('Ray Charles', 'Genius Loves Company'): dict(discs=1, copy=dict(
        date='2004', country='US', label='Concord Records / Hear Music', catno='CCD-2248-2', barcode='013431224822',
        source='Barcode on the case (checked against the ECU library catalogue record)')),
    # Spines were unreadable; details read from the case by David, Oct 2026.
    ('Fuzz(?)', 'For the Holiday(?)'): dict(discs=1, artist='Davie Allan & The Arrows', title='Fuzz for the Holidays', copy=dict(
        date='2004', label='Wicked Cool Records', catno='WC001', barcode='2534682672 (partial, as read from case)',
        note='Produced by Little Steven',
        source='Case details from David; artist confirmed by the Apple Music listing for this title')),
    ('(Spiritual Environment series)', 'Shamanic Dream'): dict(discs=1, artist='Anugama', title='Shamanic Dream', copy=dict(
        date='1990', country='West Germany', label='Meistersinger Musik', catno='NGH-CD-321',
        barcode='1831103212 (partial, as read from case)', note='Spiritual Environment series',
        source='Case details from David; Anugama album confirmed by Qobuz discography')),
    ('(artist not on spine)', 'Stormy Weather'): dict(discs=1, artist='Various', title='Stormy Weather', copy=dict(
        date='1998', country='US', label='The Isis Fund (for AT&T)',
        note='Walden Woods Project benefit concert, Wiltern Theatre, Los Angeles, April 16, 1998; produced by Don Henley; ten women singers, one standard each (Joni Mitchell sings the title track); given to AT&T customers, not sold in stores',
        source='Case details from David; confirmed by Wikipedia, Stormy Weather (AT&T album)')),
}

items = []
for r in rows('Shelf plan'):
    sec, a0, t0, discs, fmt = r['section'], r['artist'], r['title'], r['discs'], r['format']
    sk, sname = SECS[sec[0]]
    ov = OVERRIDES.get((a0, t0))
    a = ARTIST_FIX.get(a0, a0)
    if ov and 'artist' in ov: a = ov['artist']
    if sk == 'classical':
        key, a, t = CLASSICAL[t0]
        key = key.lower()
    else:
        t = ov['title'] if ov and 'title' in ov else TITLE_FIX.get(t0, t0)
        if t == 'The Beatles (White Album), 2nd copy' or (a == 'The Beatles' and t == 'The Beatles'):
            t = 'The Beatles (White Album)'
        if a == 'Grateful Dead' and t == 'The Grateful Dead': t = 'The Grateful Dead (debut)'
        if a == 'Grateful Dead' and t == 'Grateful Dead': t = 'Grateful Dead (Skull & Roses)'
        if a == 'Various' and 'Tapestry Revisited' in t: a = 'Carole King'
        key = fkey(a)
    lookup_t = 'The Beatles (White Album)' if t0 == 'The Beatles (White Album), 2nd copy' else t0
    if lookup_t == 'The Beatles (White Album)': lookup_t = 'The Beatles'
    cp = copy_info(a0, lookup_t, fmt)
    if ov:
        cp = dict(ov['copy'], **{k: v for k, v in cp.items() if k == 'found'})
        discs = ov['discs']
    items.append(dict(sec=sk, section=sname, key=key, artist=a, title=t, artist_full=a0, title_full=t0,
                      format=fmt, discs=int(discs or 1), copy=cp))

# merge identical cases into one entry with a copies list
merged = collections.OrderedDict()
for it in items:
    k = (it['sec'], it['artist'], it['title'], it['format'])
    if k in merged:
        merged[k]['copies'].append(it.pop('copy'))
    else:
        it['copies'] = [it.pop('copy')]
        merged[k] = it
items = list(merged.values())

ORDER = {
 'Beatles': ['With The Beatles','Help!','Revolver','Magical Mystery Tour','The Beatles (White Album)','Anthology 1','1','Rock ’n’ Roll','Band on the Run','GRP Artists: Songs of the Beatles'],
 'Dead1': ['The Grateful Dead (debut)','Anthem of the Sun','Aoxomoxoa','Live/Dead',"Workingman's Dead",'American Beauty','Grateful Dead (Skull & Roses)',"Europe '72","Bear's Choice",'Wake of the Flood','Skeletons From the Closet','Blues for Allah','Terrapin Station',"What a Long Strange Trip It's Been",'Shakedown Street','Reckoning','Dead Set'],
}
def by(s): return [i for i in items if i['sec'] == s]
def asort(lst): return sorted(lst, key=lambda i: (i['key'], i['title'].lower()))
def size(i): return len(i['copies'])
def split(lst, n):
    """Divide into n runs of about equal size without splitting an artist across cubbies."""
    total = sum(size(i) for i in lst); out = [[] for _ in range(n)]; idx = 0; run = 0; groups = []
    for i in lst:
        if groups and groups[-1][0]['artist'] == i['artist']: groups[-1].append(i)
        else: groups.append([i])
    target = total / n
    for g in groups:
        c = sum(size(i) for i in g)
        if idx < n - 1 and run > 0 and run + c / 2 > target * (idx + 1): idx += 1
        out[idx].extend(g); run += c
    return out

cub = {}
cub['xmas'] = ('Christmas', '', asort(by('xmas')))
cub['other'] = ('Other', 'New age, world, novelty', asort(by('other')))
cl = asort(by('classical')); h = (len(cl) + 1) // 2
cub['cl1'] = ('Classical', '', cl[:h]); cub['cl2'] = ('Classical', '', cl[h:])
j1, j2 = split(asort(by('jazz')), 2)
cub['jz1'] = ('Jazz', '', j1); cub['jz2'] = ('Jazz', '', j2)
bt = by('beatles'); bt.sort(key=lambda i: ORDER['Beatles'].index(i['title']))
cub['beatles'] = ('Beatles', 'Studio, comps, solo, tribute', bt)
gd = by('dead')
d1 = sorted([i for i in gd if i['title'] in ORDER['Dead1']], key=lambda i: ORDER['Dead1'].index(i['title']))
def drank(i):
    t = i['title']
    if i['artist'] not in ('Grateful Dead', 'Various'): return (3, i['key'], t)
    if i['artist'] == 'Various': return (4, '', t)
    if 'CD-R' in t or t in ["Soundcheck '73", 'Rainforest Benefit', 'Werewolves of Red Rocks', "09-02-68 Betty Nelson's"]: return (2, '', t)
    return (1, '', t)
rest = sorted([i for i in gd if i['title'] not in ORDER['Dead1']], key=drank)
cub['gd1'] = ('Grateful Dead', 'Studio and original-era live, in release order', d1)
cub['gd2'] = ('Grateful Dead', 'Archival, unofficial, side projects', rest)
cub['abb'] = ('Allman Brothers', '', asort(by('allmans')))
def letter(i): return i['key'][0].upper() if i['key'][0].isalpha() else '#'
for n, lst in enumerate(split(asort(by('rock')), 12)):
    rng = letter(lst[0]) + (('–' + letter(lst[-1])) if letter(lst[-1]) != letter(lst[0]) else '')
    cub['rk%d' % (n + 1)] = ('Rock & Pop ' + rng, lst[0]['artist'] + ' to ' + lst[-1]['artist'], lst)

GRID = [['xmas', 'other', None, None],
        ['cl1', 'cl2', 'jz1', 'jz2'],
        ['beatles', 'gd1', 'gd2', 'abb'],
        ['rk1', 'rk2', 'rk3', 'rk4'],
        ['rk5', 'rk6', 'rk7', 'rk8'],
        ['rk9', 'rk10', 'rk11', 'rk12']]

out_items, cubbies, seq = [], [], 0
for r, row in enumerate(GRID, 1):
    for c, cid in enumerate(row, 1):
        pos = 'R%d·C%d' % (r, c)
        if cid is None:
            cubbies.append(dict(pos=pos, name='Spare', sub='Open for growth', sec='spare', ids=[]))
            continue
        name, sub, lst = cub[cid]
        ids = []
        for it in lst:
            seq += 1
            it = dict(it); it.pop('key')
            it['id'] = seq; it['cubby'] = pos; it['cubby_name'] = name
            it['qty'] = len(it['copies'])
            dates = [x.get('date', '') for x in it['copies'] if x.get('date')]
            it['year'] = int(dates[0][:4]) if dates else None
            fixed = (it['artist_full'], it['title_full']) in OVERRIDES
            if fixed: it['artist_full'], it['title_full'] = it['artist'], it['title']
            it['check'] = not fixed and ('(?)' in it['artist_full'] + it['title_full']) or it['artist'].startswith('(') or any(x.pop('sacd_match', False) for x in it['copies'])
            if it['title'] == it['title_full']: it.pop('title_full')
            if it['artist'] == it['artist_full']: it.pop('artist_full')
            ids.append(seq); out_items.append(it)
        cubbies.append(dict(pos=pos, name=name, sub=sub, sec=lst[0]['sec'] if lst else 'spare', ids=ids))

data = dict(build=BUILD, updated=datetime.date.today().isoformat(), cubbies=cubbies, items=out_items,
            unplaced=['About 30 cases with unreadable spines', 'An unlabeled Fujifilm CD-R',
                      'A black Madama Butterfly box (second copy, or the outer box of the 3-CD Karajan set)',
                      'A Laserlight Maria Callas disc (may duplicate Live Recordings 1957-63)',
                      'A large U2 item with a photo cover'])
open(OUT, 'w', encoding='utf-8').write('window.COLLECTION=' + json.dumps(data, ensure_ascii=False) + ';\n')
left = {k: v for k, v in list(picard.items()) + list(photo.items()) if v}
print(len(out_items), 'entries,', sum(i['qty'] for i in out_items), 'cases,', sum(1 for i in out_items if i['format'] == 'SACD'), 'SACD')
print('picard matched', sum(1 for i in out_items for c in i['copies'] if 'mbid' in c))
for c in cubbies: print(c['pos'], c['name'], len(c['ids']))
